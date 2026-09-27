"""Tulis data ke SALINAN template Excel (format asli dipertahankan).

Pendekatan config-driven: mapping kolom LOG, anchor tabel inventory, dan anchor
foto per-kategori diambil dari ``template.config_json`` (fallback ke DEFAULT yang
diturunkan dari Template_BAA.xlsx). Anchor foto bisa dikalibrasi lewat Pengaturan
tanpa mengubah kode.
"""
from __future__ import annotations

import copy
from pathlib import Path

# Default mapping untuk Template_BAA.xlsx — dipakai HANYA bila template belum
# pernah di-mapping (config tanpa flag "mapped"). Setelah admin menyimpan mapping,
# config dipakai apa adanya (lihat resolve_config) tanpa digabung default ini.
DEFAULT_CONFIG = {
    "log": {
        "header_row": 2,
        "start_row": 3,
        "columns": {
            "no": "A", "lokasi": "B", "nama_barang": "C", "merk_type": "D",
            "jumlah": "E", "sn_tagging": "F", "keterangan": "G", "foto_lengkap": "H",
        },
        "sources": {},
        "value": {},
        # "first" = No/Lokasi/Kode hanya di baris pertama lokasi,
        # "repeat" = No/Lokasi/Kode diulang di setiap baris item
        "identity": "first",
        # kolom status foto per kategori: "X" bila belum ada foto, kosong bila ada
        "photo_status": {
            "dashboard": "I", "tampak_depan": "J", "teknisi": "K", "outdoor": "L",
            "indoor": "M", "sn_kit": "N", "sn_router": "O", "sn_ap": "P",
            "ping": "Q", "speed": "R", "simkopdes": "S",
        },
    },
    "detail": {
        # sel judul lokasi (opsional) — diisi nilai field data lokasi {cell: field}
        "location_cells": {},
        # sel teks statis (opsional) — diisi teks tetap {cell: text}
        "static_cells": {},
        "inventory": {
            "start_row": 32, "max_rows": 5,
            "cols": {"nama_barang": "B", "merk_type": "C", "jumlah": "E",
                     "sn_tagging": "F", "keterangan": "H"},
        },
        # anchor foto per kategori: satu sel (bila sel itu bagian dari merge) atau
        # range "B12:H26" (bila area foto tidak di-merge). Foto di-center di area.
        "photos": {
            "dashboard": "B12:H26", "tampak_depan": "B59:H73", "teknisi": "B76:H90",
            "outdoor": "K11:Q25", "indoor": "K28:Q42", "sn_kit": "K59:Q73",
            "sn_router": "K76:Q90", "sn_ap": "T9:Z23", "ping": "T28:Z42",
            "speed": "T59:Z73", "simkopdes": "T76:Z90",
        },
        # lebar foto (px); tinggi selalu 4:3 dari lebar
        "photo_w": 600,
    },
}

# Kerangka kosong untuk template yang sudah di-mapping (tanpa nilai default)
_EMPTY_CONFIG = {
    "log": {"header_row": 2, "start_row": 3, "columns": {}, "sources": {}, "value": {},
            "photo_status": {}, "identity": "first"},
    "detail": {"location_cells": {}, "static_cells": {},
               "inventory": {"start_row": 32, "max_rows": 5, "cols": {}},
               "photos": {}, "photo_w": 600},
}


def _merge(base: dict, override: dict) -> dict:
    out = copy.deepcopy(base)
    for k, v in (override or {}).items():
        if isinstance(v, dict) and isinstance(out.get(k), dict):
            out[k] = _merge(out[k], v)
        else:
            out[k] = v
    return out


def resolve_config(template_config: dict | None) -> dict:
    """Config efektif. Template yang sudah di-mapping ("mapped": true) dipakai apa
    adanya — kolom/sel yang tidak dipetakan benar-benar kosong. Template lama
    (belum pernah di-mapping) memakai DEFAULT_CONFIG."""
    tc = template_config or {}
    body = {k: v for k, v in tc.items() if k in ("log", "detail")}
    return _merge(_EMPTY_CONFIG if tc.get("mapped") else DEFAULT_CONFIG, body)


def photo_size(cfg: dict) -> tuple[int, int]:
    """(lebar, tinggi) foto — tinggi = 3/4 lebar (rasio 4:3)."""
    try:
        w = int(cfg["detail"].get("photo_w") or 600)
    except Exception:
        w = 600
    w = max(80, min(w, 4000))
    return w, round(w * 3 / 4)


def _fit(img_w: int, img_h: int, max_w: int, max_h: int) -> tuple[int, int]:
    if img_w <= 0 or img_h <= 0:
        return max_w, max_h
    r = min(max_w / img_w, max_h / img_h, 1.0)
    return int(img_w * r), int(img_h * r)


# Karakter yang dilarang Excel untuk nama worksheet
_BAD_SHEET = set(r':\/?*[]')


def _safe_sheet_title(code: str, name: str, used: set) -> str:
    """Nama worksheet = 'Kode Nama Lokasi', dibersihkan & dipotong <=31 char,
    dan dijamin unik dalam workbook."""
    raw = f"{code} {name}".strip() if name else (code or "Lokasi")
    clean = "".join(" " if ch in _BAD_SHEET else ch for ch in raw)
    clean = " ".join(clean.split()).strip("'").strip() or (code or "Lokasi")
    base = clean[:31]
    title = base
    n = 2
    while title in used or title == "":
        suffix = f" ({n})"
        title = base[:31 - len(suffix)].rstrip() + suffix
        n += 1
    used.add(title)
    return title


def _letterbox(path: str, box_w: int, box_h: int):
    """Kembalikan BytesIO PNG berukuran persis box_w x box_h: foto di-fit
    (tanpa distorsi & tanpa dipotong) lalu ditaruh di tengah kanvas putih.
    Membuat semua foto hasil export punya ukuran yang seragam."""
    import io
    from PIL import Image as PILImage
    im = PILImage.open(path)
    if im.mode not in ("RGB",):
        im = im.convert("RGB")
    im.thumbnail((box_w, box_h), PILImage.LANCZOS)   # hanya mengecilkan, jaga rasio
    canvas = PILImage.new("RGB", (box_w, box_h), (255, 255, 255))
    canvas.paste(im, ((box_w - im.width) // 2, (box_h - im.height) // 2))
    bio = io.BytesIO()
    canvas.save(bio, format="PNG")
    bio.seek(0)
    return bio


def _copy_print_settings(src, dst) -> None:
    """openpyxl copy_worksheet() tidak menyalin print area & pengaturan cetak.
    Salin manual agar sheet hasil tetap punya Print Area seperti template."""
    import copy
    try:
        if src.print_area:
            dst.print_area = src.print_area
    except Exception:
        pass
    for attr in ("page_setup", "page_margins", "print_options",
                 "sheet_properties", "print_title_rows", "print_title_cols"):
        try:
            val = getattr(src, attr)
            if val is None:
                continue
            setattr(dst, attr, copy.copy(val) if not isinstance(val, str) else val)
        except Exception:
            pass


# Perkiraan standar Excel bila lebar kolom / tinggi baris tidak diset eksplisit
_DEF_COL_CHARS = 8.43
_DEF_ROW_PT = 15.0


_MDW_CACHE: dict = {}
# Maximum Digit Width (px) perkiraan bila font tak bisa diukur langsung.
_MDW_LOOKUP = {
    "calibri": 7.0, "aptos": 7.0, "aptos narrow": 6.0, "arial": 7.0,
    "arial narrow": 6.0, "times new roman": 7.0, "verdana": 8.0,
    "tahoma": 7.0, "segoe ui": 7.0,
}


def _mdw_px(font_name: str | None, size) -> float:
    """Lebar digit maksimum (px @96dpi) untuk font+ukuran tertentu.
    Diukur langsung via PIL (akurat di mesin yang punya fontnya), fallback tabel."""
    name = (font_name or "Calibri").strip()
    try:
        size = float(size or 11.0)
    except Exception:
        size = 11.0
    key = (name.lower(), round(size, 1))
    if key in _MDW_CACHE:
        return _MDW_CACHE[key]
    val = None
    try:
        from PIL import ImageFont
        px = int(round(size * 96.0 / 72.0))
        for cand in (name, name + " Regular", name.replace(" ", ""), name.replace(" ", "") + "-Regular"):
            for ext in ("", ".ttf", ".ttc", ".otf"):
                try:
                    fnt = ImageFont.truetype(cand + ext, px)
                    val = max(fnt.getlength(str(d)) for d in range(10))
                    break
                except Exception:
                    continue
            if val:
                break
    except Exception:
        val = None
    if not val:
        val = _MDW_LOOKUP.get(key[0], 7.0)
    _MDW_CACHE[key] = val
    return val


def _col_px(ws, col_idx: int, mdw: float = 7.0) -> int:
    """Lebar kolom (1-based) dalam piksel, memperhitungkan MDW font aktif."""
    from openpyxl.utils import get_column_letter
    dim = ws.column_dimensions.get(get_column_letter(col_idx))
    w = dim.width if (dim is not None and dim.width) else None
    if not w:
        w = ws.sheet_format.defaultColWidth
        if not w:
            w = (ws.sheet_format.baseColWidth or 8) + 0.43   # perkiraan default Excel
    return int(round(w * mdw)) + 5        # rumus lebar Excel (MDW sesuai font)


def _row_px(ws, row_idx: int) -> int:
    """Tinggi baris (1-based) dalam piksel."""
    from openpyxl.utils.units import points_to_pixels
    dim = ws.row_dimensions.get(row_idx)
    h = dim.height if (dim is not None and dim.height) else None
    if not h:
        h = (ws.sheet_format.defaultRowHeight or _DEF_ROW_PT)
    return int(round(points_to_pixels(h)))


def photo_area(ws, anchor: str):
    """Area tempat foto: range eksplisit "B12:H26", atau area merge yang memuat
    sel anchor tunggal. None bila anchor satu sel yang tidak di-merge."""
    from openpyxl.worksheet.cell_range import CellRange
    anchor = (anchor or "").strip().upper()
    if not anchor:
        return None
    try:
        if ":" in anchor:
            return CellRange(anchor)
        from openpyxl.utils.cell import coordinate_to_tuple
        row, col = coordinate_to_tuple(anchor)
    except Exception:
        return None
    for m in ws.merged_cells.ranges:
        if m.min_row <= row <= m.max_row and m.min_col <= col <= m.max_col:
            return m
    return None


def _place_photo(ws, xi, anchor: str) -> None:
    """Tempel gambar di TENGAH (horizontal + vertikal) area foto: area merge yang
    memuat sel anchor, atau range eksplisit. Gambar diskalakan turun bila lebih
    besar dari area. Anchor satu sel tanpa merge -> pojok kiri-atas sel."""
    from openpyxl.utils.units import pixels_to_EMU
    from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor, AnchorMarker
    from openpyxl.drawing.xdr import XDRPositiveSize2D

    rng = photo_area(ws, anchor)
    if rng is None:                        # sel tunggal tanpa merge -> biasa
        ws.add_image(xi, anchor.split(":")[0])
        return

    try:
        f = ws.cell(row=rng.min_row, column=rng.min_col).font
        mdw = _mdw_px(f.name, f.sz)
    except Exception:
        mdw = 7.0
    area_w = sum(_col_px(ws, c, mdw) for c in range(rng.min_col, rng.max_col + 1))
    area_h = sum(_row_px(ws, r) for r in range(rng.min_row, rng.max_row + 1))
    img_w, img_h = int(xi.width), int(xi.height)
    avail_w, avail_h = max(1, area_w - 8), max(1, area_h - 8)
    if img_w > avail_w or img_h > avail_h:
        scale = min(avail_w / img_w, avail_h / img_h)
        img_w = max(1, int(img_w * scale))
        img_h = max(1, int(img_h * scale))
        xi.width, xi.height = img_w, img_h
    off_x = max(0, (area_w - img_w) // 2)
    off_y = max(0, (area_h - img_h) // 2)
    marker = AnchorMarker(col=rng.min_col - 1, colOff=pixels_to_EMU(off_x),
                          row=rng.min_row - 1, rowOff=pixels_to_EMU(off_y))
    xi.anchor = OneCellAnchor(
        _from=marker,
        ext=XDRPositiveSize2D(pixels_to_EMU(int(xi.width)), pixels_to_EMU(int(xi.height))),
    )
    ws.add_image(xi)


def _copy_row_style(ws, src_row: int, dst_row: int, max_col: int) -> None:
    """Salin format (border, font, fill, alignment, number format) & tinggi baris
    dari baris contoh template ke baris data baru — output tetap mengikuti template."""
    if dst_row == src_row:
        return
    for c in range(1, max_col + 1):
        s = ws.cell(src_row, c)
        if s.has_style:
            ws.cell(dst_row, c)._style = copy.copy(s._style)
    h = ws.row_dimensions[src_row].height
    if h:
        ws.row_dimensions[dst_row].height = h


def _has_photo(plist: list) -> bool:
    return any(p.get("path") and Path(p["path"]).exists() for p in plist or [])


def build_workbook(template_path: str, template_config: dict, locations: list[dict],
                   out_path: str, categories: list[dict] | None = None) -> dict:
    """locations: list of dict {code,name,data,inventory:[...],photos:[{category,path,...}]}
    categories: kategori foto dari Pengaturan (untuk kolom "Foto Lengkap?")."""
    import openpyxl
    from openpyxl.drawing.image import Image as XLImage

    cfg = resolve_config(template_config)
    wb = openpyxl.load_workbook(template_path)

    # Nama sheet log & detail dari config_json (disimpan saat registrasi)
    sheet_log = template_config.get("sheet_log") if template_config else None
    sheet_detail = template_config.get("sheet_detail") if template_config else None
    sheet_log = sheet_log or ("LOG" if "LOG" in wb.sheetnames else wb.sheetnames[0])
    sheet_detail = sheet_detail or (wb.sheetnames[1] if len(wb.sheetnames) > 1 else wb.sheetnames[0])

    log_ws = wb[sheet_log]
    detail_tpl = wb[sheet_detail]
    # Ganti nama sheet template detail agar salinan tak bentrok namanya
    _tpl_title = "__tpl_detail__"
    detail_tpl.title = _tpl_title

    warnings = []
    logc = cfg["log"]
    start_row = int(logc.get("start_row") or 3)
    log_row = start_row
    log_cols = logc.get("columns", {}) or {}
    log_sources = logc.get("sources", {}) or {}
    log_value = logc.get("value", {}) or {}
    log_photo = logc.get("photo_status", {}) or {}
    log_maxc = log_ws.max_column or 1
    cat_keys = [c["key"] for c in (categories or []) if c.get("key")] or list(log_photo.keys())
    pw, ph = photo_size(cfg)
    used_titles: set = set()
    _img_keep: list = []          # tahan buffer gambar sampai workbook disimpan

    def _setc(row, col, value):
        if col:
            log_ws[f"{col}{row}"] = value

    repeat_ident = (logc.get("identity") or "first") == "repeat"

    for i, loc in enumerate(locations, start=1):
        data = loc.get("data", {})
        inv = loc.get("inventory", [])
        photos = loc.get("photos", [])
        nama = data.get("nama_lokasi") or loc.get("name") or loc.get("code")
        by_cat: dict[str, list] = {}
        for p in photos:
            by_cat.setdefault(p.get("category", "uncategorized"), []).append(p)
        complete = bool(cat_keys) and all(_has_photo(by_cat.get(k)) for k in cat_keys)

        # --- baris LOG: satu baris per item inventory (list semua) ---
        rows_inv = inv if inv else [{}]          # lokasi tanpa inventory tetap 1 baris
        for j, item in enumerate(rows_inv):
            _copy_row_style(log_ws, start_row, log_row, log_maxc)
            if j == 0 or repeat_ident:           # No/Lokasi/Kode: baris pertama (atau diulang)
                _setc(log_row, log_cols.get("no"), i)
                _setc(log_row, log_cols.get("lokasi"), nama)
                _setc(log_row, log_cols.get("kode"), loc.get("code", ""))
            if j == 0:                           # kolom identitas lain hanya di baris pertama
                _setc(log_row, log_cols.get("foto_lengkap"), "SUDAH" if complete else "BELUM")
                # status foto per kategori: "X" bila belum ada, kosong bila ada
                for cat, col in log_photo.items():
                    _setc(log_row, col, "" if _has_photo(by_cat.get(cat)) else "X")
            for f in ("nama_barang", "merk_type", "jumlah", "sn_tagging", "keterangan"):
                if log_sources.get(f) != "data":
                    _setc(log_row, log_cols.get(f), item.get(f, ""))
            # kolom LOG bersumber field data lokasi — diisi di baris pertama saja
            for _f, _src in log_sources.items():
                if _src == "data" and j == 0:
                    _setc(log_row, log_cols.get(_f), data.get(log_value.get(_f, _f), ""))
            log_row += 1

        # --- sheet detail per lokasi (duplikasi template) ---
        ws = wb.copy_worksheet(detail_tpl)
        nama_only = data.get("nama_lokasi") or loc.get("name") or ""   # tanpa fallback kode
        ws.title = _safe_sheet_title(loc.get("code", f"Lokasi_{i:04d}"), nama_only, used_titles)
        _copy_print_settings(detail_tpl, ws)   # copy_worksheet tak menyalin print area/page setup

        for cell, field in cfg["detail"].get("location_cells", {}).items():
            try:
                ws[cell] = loc.get("code", "") if field == "kode" else data.get(field, "")
            except Exception:
                pass

        for cell, text in cfg["detail"].get("static_cells", {}).items():
            try:
                ws[cell] = text
            except Exception:
                pass

        invc = cfg["detail"]["inventory"]
        r0 = int(invc.get("start_row") or 1)
        max_rows = int(invc.get("max_rows") or 999)
        if len(inv) > max_rows:
            warnings.append(f"{ws.title}: {len(inv)} item inventory, hanya {max_rows} yang muat")
        for r_off, item in enumerate(inv[:max_rows]):
            r = r0 + r_off
            for field, col in (invc.get("cols") or {}).items():
                try:
                    ws[f"{col}{r}"] = item.get(field, "")
                except Exception:
                    pass

        # --- foto per kategori: tempel & center di area anchor ---
        for cat, anchor in (cfg["detail"].get("photos") or {}).items():
            plist = [p for p in by_cat.get(cat, []) if p.get("path") and Path(p["path"]).exists()]
            if not anchor or not plist:
                continue
            path = plist[0]["path"]
            try:
                try:
                    src = _letterbox(path, pw, ph)   # ukuran seragam 4:3, tanpa crop
                    xi = XLImage(src)
                    _img_keep.append(src)            # jaga buffer sampai wb.save()
                except Exception:
                    xi = XLImage(path)               # fallback: foto asli
                xi.width, xi.height = pw, ph
                _place_photo(ws, xi, anchor)
            except Exception as e:
                warnings.append(f"{ws.title}: gagal sisip foto '{cat}': {e}")

    # Hapus sheet template detail (yang sudah diganti nama) bila sudah ada salinan
    if locations and _tpl_title in wb.sheetnames and len(wb.sheetnames) > 1:
        try:
            del wb[_tpl_title]
        except Exception:
            pass
    else:
        # tak ada lokasi -> kembalikan nama sheet detail seperti semula
        wb[_tpl_title].title = sheet_detail

    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return {"path": out_path, "warnings": warnings}
