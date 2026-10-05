"""Generate output: Excel (isi salinan template) & PDF."""
from __future__ import annotations

import json
import re
import sqlite3
import threading
import time
import uuid
import zipfile
from datetime import datetime
from pathlib import Path

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import FileResponse

from .. import config, db
from ..deps import audit, current_user, get_db
from ..services import excel as excel_svc
from ..services import pdf as pdf_svc
from ..services import pdfdoc, xlsx2pdf

router = APIRouter(prefix="/api/export", tags=["export"])


class _Cancelled(Exception):
    """Dilempar dari callback progress saat user membatalkan export."""


def _filter_where(q: str, status: str, creator: int, date: str, user=None, wil: str = ""):
    """Bangun klausa WHERE yang sama dengan GET /api/locations (Log Lokasi).
    Read-all: semua user boleh mengekspor lokasi mana pun."""
    where, params = ["deleted_at IS NULL"], []
    if q.strip():
        like = f"%{q.strip()}%"
        where.append("(code LIKE ? OR name LIKE ? OR old_codes LIKE ?)")
        params += [like, like, like]
    if status in ("draft", "selesai"):
        where.append("status = ?")
        params.append(status)
    if creator:
        where.append("created_by = ?")
        params.append(creator)
    if date.strip():
        where.append("date(created_at,'localtime') = ?")
        params.append(date.strip())
    if wil.strip():
        from .locations import wil_where
        w, p = wil_where(wil)
        where.append(w); params += p
    wsql = ("WHERE " + " AND ".join(where)) if where else ""
    return wsql, params


def _gather_locations(conn: sqlite3.Connection, scope: str, loc_id: int | None,
                      q: str = "", status: str = "all", creator: int = 0,
                      date: str = "", user=None, wil: str = "") -> list[dict]:
    if scope == "all":
        rows = conn.execute("SELECT * FROM locations WHERE deleted_at IS NULL ORDER BY id").fetchall()
    elif scope == "filter":
        wsql, params = _filter_where(q, status, creator, date, user, wil)
        rows = conn.execute(f"SELECT * FROM locations {wsql} ORDER BY id", params).fetchall()
    else:
        if not loc_id:
            raise HTTPException(400, "loc_id is required for scope 'one'")
        rows = conn.execute("SELECT * FROM locations WHERE id=? AND deleted_at IS NULL", (loc_id,)).fetchall()
    if not rows:
        raise HTTPException(404, "No locations to export")
    out = []
    for r in rows:
        inv = conn.execute(
            "SELECT nama_barang,merk_type,jumlah,sn_tagging,keterangan FROM inventory_items "
            "WHERE location_id=? ORDER BY sort_order,id", (r["id"],)).fetchall()
        photos = conn.execute(
            "SELECT category,path,ocr_serial FROM photos WHERE location_id=? ORDER BY id",
            (r["id"],)).fetchall()
        data = json.loads(r["data_json"])
        # Wilayah sesuai tulisan teknisi -> bisa dipetakan ke kolom Log / sel Detail seperti field data
        for k in ("wil_desa", "wil_kec", "wil_kab", "wil_prov", "wil_kode"):
            data[k] = r[k] or ""
        out.append({
            "id": r["id"], "code": r["code"], "name": r["name"],
            "data": data,
            "inventory": [dict(i) for i in inv],
            "photos": [{**dict(p), "path": str(db.fpath(conn, p["path"]))} for p in photos],
            "scan": (str(db.fpath(conn, sc["path"])) if (sc := conn.execute(
                "SELECT path FROM scan_docs WHERE location_id=?", (r["id"],)).fetchone()) else None),
        })
    return out


def _stamp(prefix: str, ext: str) -> Path:
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    return config.OUTPUT_DIR / f"{prefix}_{ts}.{ext}"


def _safe_name(s: str) -> str:
    """Nama file aman (untuk entri ZIP)."""
    s = re.sub(r'[\\/:*?"<>|]+', " ", (s or "")).strip()
    return re.sub(r"\s+", " ", s) or "lokasi"


def _tpl_row(conn: sqlite3.Connection) -> dict | None:
    """Template aktif (path sudah absolut), atau None."""
    t = conn.execute("SELECT * FROM templates WHERE active=1 ORDER BY id DESC LIMIT 1").fetchone()
    return {**dict(t), "path": str(db.fpath(conn, t["path"]))} if t else None


def _wil_folder(loc: dict) -> str:
    """Folder ZIP per wilayah: 'Provinsi/Kab-Kota/' memakai nama resmi bila ada kode (agar ejaan
    berbeda tetap satu folder); isian manual memakai tulisannya; tanpa wilayah -> 'Tanpa wilayah/'."""
    from ..services import wilayah
    d = loc.get("data") or {}
    k = d.get("wil_kode") or ""
    if k and wilayah.name_of(k[:5]):
        prov, kab = wilayah.name_of(k[:2]), wilayah.short_kab(wilayah.name_of(k[:5]))
    elif d.get("wil_prov") or d.get("wil_kab"):
        prov, kab = d.get("wil_prov") or "-", d.get("wil_kab") or "-"
    else:
        return "Tanpa wilayah/"
    return f"{_safe_name(prov)}/{_safe_name(kab)}/"


def _active_template(conn: sqlite3.Connection):
    """Template aktif + config, atau (None, None) bila tak ada / file hilang."""
    tpl = _tpl_row(conn)
    if not tpl or not Path(tpl["path"]).exists():
        return None, None
    tcfg = json.loads(tpl["config_json"]) if tpl["config_json"] else {}
    tcfg["sheet_log"] = tpl["sheet_log"]
    tcfg["sheet_detail"] = tpl["sheet_detail"]
    return tpl, tcfg


def _photo_anchors(tcfg: dict) -> list[str]:
    """Anchor/area foto efektif (mengikuti config ter-mapping atau default)."""
    return [a for a in (excel_svc.resolve_config(tcfg)["detail"].get("photos") or {}).values() if a]


def _template_pdf(tpl, tcfg, locs, out_pdf: Path, cats=None, progress=None, stage=None) -> tuple[bool, str]:
    """Isi template -> xlsx (detail saja, tanpa LOG) -> konversi PDF (Excel/LO).
    (True, peringatan_kertas) bila berhasil; False agar pemanggil fallback ke reportlab.
    Peringatan terisi bila ukuran halaman PDF tidak sama dengan Page Setup template
    (biasanya karena printer default Windows tidak mendukung ukuran kertas itu)."""
    if not xlsx2pdf.available():
        return False, ""
    import openpyxl
    tmp_xlsx = out_pdf.with_suffix(".xlsx")
    ok = False
    warn = ""
    try:
        excel_svc.build_workbook(tpl["path"], tcfg, locs, str(tmp_xlsx), cats, progress)
        if stage:
            stage("conv")
        wb = openpyxl.load_workbook(str(tmp_xlsx))
        log_name = tcfg.get("sheet_log")
        if log_name and log_name in wb.sheetnames and len(wb.sheetnames) > 1:
            del wb[log_name]                  # PDF hanya halaman detail lokasi
        wb.save(str(tmp_xlsx))
        anchors = _photo_anchors(tcfg)
        ok = xlsx2pdf.xlsx_to_pdf(str(tmp_xlsx), str(out_pdf), anchors)
        if ok:
            warn = xlsx2pdf.paper_warning(str(tmp_xlsx), str(out_pdf))
    except _Cancelled:
        raise
    except Exception:
        ok = False
    finally:
        try:
            tmp_xlsx.unlink()
        except OSError:
            pass
    return bool(ok) and out_pdf.exists(), warn


def _render_pdf(conn, locs, out_pdf: Path, cats, title) -> tuple[str, str]:
    """PDF mengikuti template aktif bila memungkinkan; jika tidak, pakai
    generator bawaan (reportlab). Mengembalikan ('template'|'builtin', peringatan_kertas)."""
    tpl, tcfg = _active_template(conn)
    return _render_pdf_with(tpl, tcfg, locs, out_pdf, cats, title)


def _render_pdf_with(tpl, tcfg, locs, out_pdf: Path, cats, title, progress=None, stage=None) -> tuple[str, str]:
    """Seperti _render_body, plus scan PDF tiap lokasi di depan halaman BAA-nya.
    Lokasi tanpa scan tetap diekspor apa adanya."""
    if not any(loc.get("scan") for loc in locs):
        return _render_body(tpl, tcfg, locs, out_pdf, cats, title, progress, stage)
    parts, tmps, notes = [], [], []
    mode, warn, n = "template", "", len(locs)
    try:
        for i, loc in enumerate(locs, start=1):
            if progress and n > 1:
                progress(i, n)
            body = out_pdf.with_name(f"{out_pdf.stem}__b{i}.pdf")
            m, w = _render_body(tpl, tcfg, [loc], body, cats, title, None, stage)
            tmps.append(body)
            mode, warn = (m if i == 1 else mode), (warn or w)
            if loc.get("scan") and Path(loc["scan"]).exists():
                parts.append(loc["scan"])
                if not notes:
                    notes += _scan_paper_note(loc["scan"], body)
            parts.append(body)
        skipped = pdfdoc.concat(parts, out_pdf)
    finally:
        for t in tmps:
            try:
                t.unlink()
            except OSError:
                pass
    if skipped:
        notes.append(f"{len(skipped)} scan PDF could not be read and was left out.")
    return mode, " ".join([x for x in (warn, *notes) if x])


def _scan_paper_note(scan: str, body: Path) -> list[str]:
    a, b = pdfdoc.first_page_size(scan), pdfdoc.first_page_size(body)
    if not a or not b:
        return []
    same = lambda x, y: abs(x[0] - y[0]) < 6 and abs(x[1] - y[1]) < 6
    if same(a, b) or same(a, (b[1], b[0])):
        return []
    return [f"Scan PDF is {xlsx2pdf._paper_name(*a)} while the BAA pages are {xlsx2pdf._paper_name(*b)}."]


def _render_body(tpl, tcfg, locs, out_pdf: Path, cats, title, progress=None, stage=None) -> tuple[str, str]:
    if tpl:
        ok, warn = _template_pdf(tpl, tcfg, locs, out_pdf, cats, progress, stage)
        if ok:
            return "template", warn
    pdf_svc.build_pdf(locs, cats, str(out_pdf), app_title=title)
    return "builtin", ""


def _log_export(conn, locs, kind: str, user) -> None:
    now = db.now_iso()
    conn.executemany("INSERT INTO export_log(location_id, kind, user_id, created_at) VALUES(?,?,?,?)",
                     [(loc["id"], kind, user["id"], now) for loc in locs if loc.get("id")])
    conn.commit()


@router.get("/excel")
def export_excel(scope: str = Query("one"), loc_id: int | None = None,
                 q: str = Query(""), status: str = Query("all"),
                 creator: int = Query(0), date: str = Query(""), wil: str = Query(""),
                 conn: sqlite3.Connection = Depends(get_db), user=Depends(current_user)):
    tpl = _tpl_row(conn)
    if not tpl:
        raise HTTPException(400, "No active template. Add one in the Template menu first.")
    if not Path(tpl["path"]).exists():
        raise HTTPException(400, "Template file is missing on the server.")
    locs = _gather_locations(conn, scope, loc_id, q, status, creator, date, user, wil)
    tcfg = json.loads(tpl["config_json"]) if tpl["config_json"] else {}
    tcfg["sheet_log"] = tpl["sheet_log"]
    tcfg["sheet_detail"] = tpl["sheet_detail"]
    prefix = locs[0]["code"] if scope == "one" else "Log_BAA"
    out = _stamp(prefix, "xlsx")
    cats = db.get_setting(conn, "photo_categories", [])
    res = excel_svc.build_workbook(tpl["path"], tcfg, locs, str(out), cats)
    # Pusatkan foto secara akurat via Excel (bila tersedia) — bebas font/render
    try:
        anchors = _photo_anchors(tcfg)
        xlsx2pdf.recenter_images_excel(str(out), anchors)
    except Exception:
        pass
    _log_export(conn, locs, "excel", user)
    audit(conn, user, "export_excel", "export", scope, out.name)
    resp = FileResponse(str(out), filename=out.name,
                        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    if res.get("warnings"):
        resp.headers["X-Export-Warnings"] = str(len(res["warnings"]))
    return resp


@router.get("/pdf")
def export_pdf(scope: str = Query("one"), loc_id: int | None = None,
               q: str = Query(""), status: str = Query("all"),
               creator: int = Query(0), date: str = Query(""), wil: str = Query(""),
               conn: sqlite3.Connection = Depends(get_db), user=Depends(current_user)):
    locs = _gather_locations(conn, scope, loc_id, q, status, creator, date, user, wil)
    cats = db.get_setting(conn, "photo_categories", [])
    title = db.app_title(conn)
    out = _stamp(locs[0]["code"] if scope == "one" else "BAA", "pdf")
    mode, warn = _render_pdf(conn, locs, out, cats, title)
    _log_export(conn, locs, "pdf", user)
    audit(conn, user, "export_pdf", "export", scope, out.name)
    resp = FileResponse(str(out), filename=out.name, media_type="application/pdf")
    resp.headers["X-PDF-Mode"] = mode
    if warn:
        resp.headers["X-Paper-Warning"] = warn
    return resp


@router.get("/pdf-zip")
def export_pdf_zip(scope: str = Query("filter"), loc_id: int | None = None,
                   q: str = Query(""), status: str = Query("all"),
                   creator: int = Query(0), date: str = Query(""), wil: str = Query(""),
                   conn: sqlite3.Connection = Depends(get_db), user=Depends(current_user)):
    """Satu PDF detail per lokasi (tanpa LOG sheet), dibundel dalam satu ZIP."""
    locs = _gather_locations(conn, scope, loc_id, q, status, creator, date, user, wil)
    cats = db.get_setting(conn, "photo_categories", [])
    title = db.app_title(conn)
    zpath = _stamp("PDF_BAA", "zip")
    used: set[str] = set()
    paper_warn = ""
    with zipfile.ZipFile(zpath, "w", zipfile.ZIP_DEFLATED) as zf:
        for loc in locs:
            nm = _safe_name(f"{loc.get('code','')} {loc.get('name','') or loc.get('data',{}).get('nama_lokasi','')}")
            entry = f"{nm}.pdf"
            i = 2
            while entry.lower() in used:
                entry = f"{nm} ({i}).pdf"; i += 1
            used.add(entry.lower())
            tmp = config.OUTPUT_DIR / f"_tmp_{loc['id']}.pdf"
            _, w = _render_pdf(conn, [loc], tmp, cats, title)
            paper_warn = paper_warn or w
            zf.write(str(tmp), entry)
            try:
                tmp.unlink()
            except OSError:
                pass
    _log_export(conn, locs, "pdfzip", user)
    audit(conn, user, "export_pdf_zip", "export", scope, zpath.name)
    resp = FileResponse(str(zpath), filename=zpath.name, media_type="application/zip")
    if paper_warn:
        resp.headers["X-Paper-Warning"] = paper_warn
    return resp


# ================= Export sebagai pekerjaan latar (progress + batal) =================
# Alur: POST /jobs -> {id} ; GET /jobs/{id} (poll tahap & progres) ; GET /jobs/{id}/file ; DELETE /jobs/{id} (batal).
# Validasi (template aktif, lokasi ada) tetap sinkron di POST supaya error langsung terlihat.
_JOBS: dict[str, dict] = {}
_JOBS_LOCK = threading.Lock()
_JOB_TTL = 3600


def _job_public(j: dict) -> dict:
    return {k: j[k] for k in ("id", "kind", "name", "stage", "done", "total", "state", "error", "warn", "img_warn",
                              "scan_missing")}


def _purge_jobs() -> None:
    now = time.time()
    with _JOBS_LOCK:
        for k in [k for k, j in _JOBS.items() if now - j["created"] > _JOB_TTL]:
            _JOBS.pop(k, None)


def _run_job(j: dict, locs: list, tpl, tcfg, cats, title) -> None:
    def prog(i, n):
        if j["cancel"]:
            raise _Cancelled()
        if n > 1:                                        # 1 lokasi = progres tak berangka
            j["done"], j["total"] = i - 1, n

    def stage(s):
        if j["cancel"]:
            raise _Cancelled()
        j["stage"] = s
    try:
        kind = j["kind"]
        if kind == "excel":
            out = _stamp(locs[0]["code"] if j["scope"] == "one" else "Log_BAA", "xlsx")
            stage("fill")
            res = excel_svc.build_workbook(tpl["path"], tcfg, locs, str(out), cats, prog)
            j["done"] = j["total"]
            stage("photo")
            try:
                xlsx2pdf.recenter_images_excel(str(out), _photo_anchors(tcfg))
            except Exception:
                pass
            j["img_warn"] = len(res.get("warnings") or [])
            j.update(path=str(out), filename=out.name,
                     media="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        elif kind == "pdf":
            out = _stamp(locs[0]["code"] if j["scope"] == "one" else "BAA", "pdf")
            stage("fill")
            _, warn = _render_pdf_with(tpl, tcfg, locs, out, cats, title, prog, stage)
            j["warn"] = warn
            j.update(path=str(out), filename=out.name, media="application/pdf")
        else:                                            # pdfzip
            zpath = _stamp("PDF_BAA", "zip")
            stage("conv")
            used: set[str] = set()
            with zipfile.ZipFile(zpath, "w", zipfile.ZIP_DEFLATED) as zf:
                for n, loc in enumerate(locs, start=1):
                    if j["cancel"]:
                        raise _Cancelled()
                    j["done"], j["total"] = n - 1, len(locs)
                    nm = _safe_name(f"{loc.get('code','')} {loc.get('name','') or loc.get('data',{}).get('nama_lokasi','')}")
                    folder = _wil_folder(loc) if j.get("group") == "wilayah" else ""
                    entry = f"{folder}{nm}.pdf"
                    k = 2
                    while entry.lower() in used:
                        entry = f"{folder}{nm} ({k}).pdf"; k += 1
                    used.add(entry.lower())
                    tmp = config.OUTPUT_DIR / f"_tmp_{j['id']}_{loc['id']}.pdf"
                    _, w = _render_pdf_with(tpl, tcfg, [loc], tmp, cats, title)
                    j["warn"] = j["warn"] or w
                    zf.write(str(tmp), entry)
                    try:
                        tmp.unlink()
                    except OSError:
                        pass
                j["done"] = len(locs)
                stage("zip")
            j.update(path=str(zpath), filename=zpath.name, media="application/zip")
        j["stage"], j["state"] = "dl", "ok"
    except _Cancelled:
        j["state"] = "cancel"
    except Exception as e:                              # pesan singkat untuk kartu progress
        j["state"], j["error"] = "err", (str(e) or e.__class__.__name__)[:200]


@router.post("/jobs")
def start_job(kind: str = Query(...), scope: str = Query("one"), loc_id: int | None = None,
              q: str = Query(""), status: str = Query("all"), creator: int = Query(0), date: str = Query(""), wil: str = Query(""),
              group: str = Query(""),
              conn: sqlite3.Connection = Depends(get_db), user=Depends(current_user)):
    if kind not in ("excel", "pdf", "pdfzip"):
        raise HTTPException(400, "Unknown export type")
    _purge_jobs()
    locs = _gather_locations(conn, scope, loc_id, q, status, creator, date, user, wil)
    tpl, tcfg = _active_template(conn)
    if kind == "excel":
        t = conn.execute("SELECT * FROM templates WHERE active=1 ORDER BY id DESC LIMIT 1").fetchone()
        if not t:
            raise HTTPException(400, "No active template. Add one in the Template menu first.")
        if not tpl:
            raise HTTPException(400, "Template file is missing on the server.")
    cats = db.get_setting(conn, "photo_categories", [])
    title = db.app_title(conn)
    n = len(locs)
    one = f"{locs[0]['code']}" + (f" · {locs[0].get('name')}" if locs[0].get("name") else "")
    name = {"excel": (one + " · Excel") if scope == "one" else f"Location Log · {n} location{'s' if n != 1 else ''}",
            "pdf": (one + " · PDF") if scope == "one" else f"BAA PDF · {n} locations",
            "pdfzip": f"PDF per location · {n} location{'s' if n != 1 else ''}"}[kind]
    jid = uuid.uuid4().hex[:12]
    j = {"id": jid, "user_id": user["id"], "kind": kind, "scope": scope, "name": name, "stage": "queued",
         "done": 0, "total": n if n > 1 else 0, "state": "run", "error": "", "warn": "", "img_warn": 0,
         "scan_missing": sum(1 for loc in locs if not loc.get("scan")),
         "cancel": False, "created": time.time(), "path": None, "filename": None, "media": None,
         "group": group if kind == "pdfzip" else ""}
    with _JOBS_LOCK:
        _JOBS[jid] = j
    _log_export(conn, locs, kind, user)
    audit(conn, user, {"excel": "export_excel", "pdf": "export_pdf", "pdfzip": "export_pdf_zip"}[kind], "export", scope, name)
    threading.Thread(target=_run_job, args=(j, locs, tpl, tcfg, cats, title), daemon=True).start()
    return _job_public(j)


def _own_job(jid: str, user) -> dict:
    j = _JOBS.get(jid)
    if not j or j["user_id"] != user["id"]:
        raise HTTPException(404, "Export not found")
    return j


@router.get("/jobs/{jid}")
def job_status(jid: str, user=Depends(current_user)):
    return _job_public(_own_job(jid, user))


@router.delete("/jobs/{jid}")
def job_cancel(jid: str, user=Depends(current_user)):
    j = _own_job(jid, user)
    j["cancel"] = True
    return _job_public(j)


@router.get("/jobs/{jid}/file")
def job_file(jid: str, user=Depends(current_user)):
    j = _own_job(jid, user)
    if j["state"] != "ok" or not j["path"] or not Path(j["path"]).exists():
        raise HTTPException(409, "Export is not ready")
    resp = FileResponse(j["path"], filename=j["filename"], media_type=j["media"])
    if j["img_warn"]:
        resp.headers["X-Export-Warnings"] = str(j["img_warn"])
    if j["warn"]:
        resp.headers["X-Paper-Warning"] = j["warn"]
    if j["scan_missing"]:
        resp.headers["X-Scan-Missing"] = str(j["scan_missing"])
    return resp
