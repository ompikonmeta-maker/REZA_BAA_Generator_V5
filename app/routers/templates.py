"""Template Excel: upload, deteksi worksheet, registrasi (pilih log & detail)."""
from __future__ import annotations

import sqlite3
import uuid
from pathlib import Path

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from pydantic import BaseModel

from .. import config, db
from ..deps import audit, current_user, get_db, require_admin

router = APIRouter(prefix="/api/templates", tags=["templates"])


def _inspect_workbook(path: Path, preview_rows: int = 300, preview_cols: int = 78) -> list[dict]:
    """Kembalikan tiap worksheet: sel non-kosong (koordinat+nilai) + merged ranges.

    Dipakai UI Template Mapping untuk memilih anchor sel per kategori foto,
    kolom LOG, dan anchor tabel inventory.
    """
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    sheets = []
    for ws in wb.worksheets:
        cells = []
        maxr = min(ws.max_row or 0, preview_rows)
        maxc = min(ws.max_column or 0, preview_cols)
        for r in range(1, maxr + 1):
            for c in range(1, maxc + 1):
                v = ws.cell(r, c).value
                if v is not None:
                    cells.append({"ref": ws.cell(r, c).coordinate, "text": str(v)[:120]})
        merged = [str(rng) for rng in ws.merged_cells.ranges]
        sheets.append({"name": ws.title, "max_row": ws.max_row or 0,
                       "max_col": ws.max_column or 0, "cells": cells, "merged": merged})
    wb.close()
    return sheets


def _sheet_render(ws, max_rows: int = 300, max_cols: int = 78) -> dict:
    """Data untuk menggambar worksheet di UI (Pick on sheet): lebar kolom &
    tinggi baris (px), merge, border, tebal, rata, ukuran font, warna isi, teks."""
    from openpyxl.utils import get_column_letter
    nr = min(ws.max_row or 1, max_rows)
    nc = min(ws.max_column or 1, max_cols)
    dw = ws.sheet_format.defaultColWidth or 8.43
    dh = ws.sheet_format.defaultRowHeight or 15
    cols = []
    for c in range(1, nc + 1):
        d = ws.column_dimensions.get(get_column_letter(c))
        w = d.width if (d is not None and d.width) else dw
        cols.append(0 if (d is not None and d.hidden) else int(round(w * 7 + 5)))
    rows = []
    for r in range(1, nr + 1):
        d = ws.row_dimensions.get(r)
        h = d.height if (d is not None and d.height) else dh
        rows.append(0 if (d is not None and d.hidden) else int(round(h * 4 / 3)))
    cells = {}
    for row in ws.iter_rows(min_row=1, max_row=nr, max_col=nc):
        for c in row:
            info = {}
            if c.value is not None:
                info["v"] = str(c.value)[:200]
            b = c.border
            bd = "".join(k for k, sd in (("t", b.top), ("r", b.right), ("b", b.bottom), ("l", b.left))
                         if sd is not None and sd.style)
            if bd:
                info["bd"] = bd
            try:
                if c.fill is not None and c.fill.fill_type == "solid":
                    rgb = c.fill.fgColor.rgb
                    if isinstance(rgb, str) and len(rgb) >= 6 and rgb not in ("00000000",):
                        info["f"] = "#" + rgb[-6:]
            except Exception:
                pass
            if c.font is not None:
                if c.font.b:
                    info["b"] = 1
                if c.font.sz and float(c.font.sz) != 11:
                    info["sz"] = float(c.font.sz)
            a = c.alignment
            if a is not None:
                if a.horizontal:
                    info["h"] = a.horizontal
                if a.vertical:
                    info["va"] = a.vertical
                if a.text_rotation:
                    info["rot"] = int(a.text_rotation)
                if a.wrap_text:
                    info["wr"] = 1
            if info:
                cells[c.coordinate] = info
    return {"name": ws.title, "cols": cols, "rows": rows, "cells": cells,
            "merged": [str(m) for m in ws.merged_cells.ranges]}


@router.post("/inspect")
async def inspect_template(file: UploadFile = File(...),
                           conn: sqlite3.Connection = Depends(get_db),
                           user=Depends(require_admin)):
    """Simpan sementara & kembalikan daftar worksheet + preview untuk dipilih."""
    if not (file.filename or "").lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(400, "Must be an .xlsx/.xlsm file")
    tmp_name = f"pending_{uuid.uuid4().hex}.xlsx"
    tmp_path = db.templates_dir(conn) / tmp_name
    tmp_path.write_bytes(await file.read())
    try:
        sheets = _inspect_workbook(tmp_path)
    except Exception as e:
        tmp_path.unlink(missing_ok=True)
        raise HTTPException(400, f"Can't read Excel: {e}")
    return {"pending_file": tmp_name, "orig_name": file.filename, "sheets": sheets}


class RegisterIn(BaseModel):
    pending_file: str
    name: str
    orig_name: str = "template.xlsx"
    sheet_log: str
    sheet_detail: str
    config: dict = {}
    activate: bool = True


@router.post("/register")
def register_template(body: RegisterIn, conn: sqlite3.Connection = Depends(get_db),
                      user=Depends(require_admin)):
    import json
    if body.sheet_log == body.sheet_detail:
        raise HTTPException(400, "Log and Detail must be different sheets")
    if "/" in body.pending_file or "\\" in body.pending_file:
        raise HTTPException(400, "Invalid file name")
    src = db.templates_dir(conn) / body.pending_file
    if not src.exists():
        raise HTTPException(400, "Upload expired — please upload again")
    final_name = f"tpl_{uuid.uuid4().hex}.xlsx"
    final_path = db.templates_dir(conn) / final_name
    src.rename(final_path)
    now = db.now_iso()
    # Template pertama (belum ada yang aktif) langsung diaktifkan
    activate = body.activate or not conn.execute(
        "SELECT 1 FROM templates WHERE active=1").fetchone()
    cur = conn.execute(
        "INSERT INTO templates(name,filename,path,sheet_log,sheet_detail,config_json,"
        "active,uploaded_by,created_at) VALUES(?,?,?,?,?,?,?,?,?)",
        (body.name, body.orig_name, db.rel_path(conn, final_path), body.sheet_log, body.sheet_detail,
         json.dumps(body.config, ensure_ascii=False), 1 if activate else 0,
         user["id"], now),
    )
    if activate:
        conn.execute("UPDATE templates SET active=0 WHERE id<>?", (cur.lastrowid,))
    conn.commit()
    audit(conn, user, "register", "template", cur.lastrowid, body.name)
    return {"ok": True, "id": cur.lastrowid, "active": bool(activate)}


@router.get("")
def list_templates(conn: sqlite3.Connection = Depends(get_db), user=Depends(current_user)):
    import json
    rows = conn.execute(
        "SELECT id,name,filename,sheet_log,sheet_detail,active,created_at,config_json "
        "FROM templates ORDER BY id DESC"
    ).fetchall()
    out = []
    for r in rows:
        d = dict(r)
        try:
            d["mapped"] = bool(json.loads(d.pop("config_json") or "{}").get("mapped"))
        except Exception:
            d["mapped"] = False
        out.append(d)
    return out


@router.post("/{tpl_id}/activate")
def activate_template(tpl_id: int, conn: sqlite3.Connection = Depends(get_db),
                      user=Depends(require_admin)):
    if not conn.execute("SELECT 1 FROM templates WHERE id=?", (tpl_id,)).fetchone():
        raise HTTPException(404, "Template not found")
    conn.execute("UPDATE templates SET active=CASE WHEN id=? THEN 1 ELSE 0 END", (tpl_id,))
    conn.commit()
    audit(conn, user, "activate", "template", tpl_id)
    return {"ok": True}


class RenameIn(BaseModel):
    name: str


@router.patch("/{tpl_id}")
def rename_template(tpl_id: int, body: RenameIn, conn: sqlite3.Connection = Depends(get_db),
                    user=Depends(require_admin)):
    """Ganti nama template terdaftar."""
    name = body.name.strip()
    if not name:
        raise HTTPException(400, "Name can't be empty")
    if not conn.execute("SELECT 1 FROM templates WHERE id=?", (tpl_id,)).fetchone():
        raise HTTPException(404, "Template not found")
    conn.execute("UPDATE templates SET name=? WHERE id=?", (name, tpl_id))
    conn.commit()
    audit(conn, user, "rename", "template", tpl_id, name)
    return {"ok": True, "name": name}


@router.delete("/{tpl_id}")
def delete_template(tpl_id: int, conn: sqlite3.Connection = Depends(get_db),
                    user=Depends(require_admin)):
    """Hapus template beserta file-nya. Template aktif tidak boleh dihapus."""
    row = conn.execute("SELECT * FROM templates WHERE id=?", (tpl_id,)).fetchone()
    if not row:
        raise HTTPException(404, "Template not found")
    if row["active"]:
        raise HTTPException(400, "Can't delete the active template — activate another one first")
    try:
        db.fpath(conn, row["path"]).unlink(missing_ok=True)
    except Exception:
        pass
    conn.execute("DELETE FROM templates WHERE id=?", (tpl_id,))
    conn.commit()
    audit(conn, user, "delete", "template", tpl_id, row["name"])
    return {"ok": True}


@router.get("/{tpl_id}")
def get_template(tpl_id: int, conn: sqlite3.Connection = Depends(get_db),
                 user=Depends(current_user)):
    import json
    row = conn.execute("SELECT * FROM templates WHERE id=?", (tpl_id,)).fetchone()
    if not row:
        raise HTTPException(404, "Template not found")
    d = dict(row)
    d["config"] = json.loads(row["config_json"]) if row["config_json"] else {}
    d.pop("config_json", None)
    return d


@router.get("/{tpl_id}/sheets")
def template_sheets(tpl_id: int, conn: sqlite3.Connection = Depends(get_db),
                    user=Depends(require_admin)):
    """Baca ulang worksheet template terdaftar (untuk UI mapping)."""
    row = conn.execute("SELECT * FROM templates WHERE id=?", (tpl_id,)).fetchone()
    if not row:
        raise HTTPException(404, "Template not found")
    if not db.fpath(conn, row["path"]).exists():
        raise HTTPException(400, "Template file is missing on the server")
    return {"sheets": _inspect_workbook(db.fpath(conn, row["path"]))}


class MappingIn(BaseModel):
    config: dict            # {log:{...}, detail:{...}} anchor mapping
    sheet_log: str | None = None
    sheet_detail: str | None = None


@router.put("/{tpl_id}/mapping")
def save_mapping(tpl_id: int, body: MappingIn, conn: sqlite3.Connection = Depends(get_db),
                 user=Depends(require_admin)):
    import json
    row = conn.execute("SELECT * FROM templates WHERE id=?", (tpl_id,)).fetchone()
    if not row:
        raise HTTPException(404, "Template not found")
    if body.sheet_log and body.sheet_log == body.sheet_detail:
        raise HTTPException(400, "Log and Detail must be different sheets")
    fields = ["config_json=?"]
    params: list = [json.dumps(body.config, ensure_ascii=False)]
    if body.sheet_log:
        fields.append("sheet_log=?"); params.append(body.sheet_log)
    if body.sheet_detail:
        fields.append("sheet_detail=?"); params.append(body.sheet_detail)
    params.append(tpl_id)
    conn.execute(f"UPDATE templates SET {','.join(fields)} WHERE id=?", params)
    conn.commit()
    audit(conn, user, "save_mapping", "template", tpl_id)
    return {"ok": True}


@router.get("/{tpl_id}/render")
def template_render(tpl_id: int, sheet: str, conn: sqlite3.Connection = Depends(get_db),
                    user=Depends(require_admin)):
    """Gambar worksheet (lebar/tinggi, merge, border, teks) untuk Pick on sheet."""
    import openpyxl
    row = conn.execute("SELECT path FROM templates WHERE id=?", (tpl_id,)).fetchone()
    if not row:
        raise HTTPException(404, "Template not found")
    if not db.fpath(conn, row["path"]).exists():
        raise HTTPException(400, "Template file is missing on the server")
    wb = openpyxl.load_workbook(db.fpath(conn, row["path"]))
    try:
        if sheet not in wb.sheetnames:
            raise HTTPException(404, "Worksheet not found")
        return _sheet_render(wb[sheet])
    finally:
        wb.close()


class TrialIn(BaseModel):
    config: dict
    sheet_log: str
    sheet_detail: str
    limit: int = 2


def _trial_build(tpl_id: int, body: TrialIn, conn: sqlite3.Connection) -> tuple[Path, dict, dict]:
    """Bangun workbook uji dari mapping (belum disimpan) + lokasi terbaru."""
    from ..routers.export import _gather_locations
    from ..services import excel as excel_svc
    row = conn.execute("SELECT path FROM templates WHERE id=?", (tpl_id,)).fetchone()
    if not row:
        raise HTTPException(404, "Template not found")
    if not db.fpath(conn, row["path"]).exists():
        raise HTTPException(400, "Template file is missing on the server")
    ids = [r["id"] for r in conn.execute(
        "SELECT id FROM locations WHERE deleted_at IS NULL ORDER BY id DESC LIMIT ?",
        (max(1, min(body.limit, 10)),)).fetchall()]
    if not ids:
        raise HTTPException(400, "No locations yet to preview")
    locs = []
    for i in reversed(ids):
        locs += _gather_locations(conn, "one", i)
    tcfg = dict(body.config or {})
    tcfg["mapped"] = True
    tcfg["sheet_log"] = body.sheet_log
    tcfg["sheet_detail"] = body.sheet_detail
    cats = db.get_setting(conn, "photo_categories", [])
    out = config.OUTPUT_DIR / f"_trial_{uuid.uuid4().hex}.xlsx"
    res = excel_svc.build_workbook(str(db.fpath(conn, row["path"])), tcfg, locs, str(out), cats)
    return out, res, excel_svc.resolve_config(tcfg)


@router.post("/{tpl_id}/preview")
def template_preview(tpl_id: int, body: TrialIn, conn: sqlite3.Connection = Depends(get_db),
                     user=Depends(require_admin)):
    """Sheet LOG hasil mapping yang sedang diedit (tanpa menyimpan), digambar
    seperti Excel: header asli template + baris data lokasi terbaru."""
    import openpyxl
    out, res, cfg = _trial_build(tpl_id, body, conn)
    try:
        wb = openpyxl.load_workbook(str(out))
        ws = wb[body.sheet_log]
        hr = int(cfg["log"].get("header_row") or 2)
        sr = int(cfg["log"].get("start_row") or hr + 1)
        last = sr
        for r in range(sr, min(ws.max_row or sr, sr + 80) + 1):
            if any(ws.cell(r, c).value not in (None, "") for c in range(1, (ws.max_column or 1) + 1)):
                last = r
        render = _sheet_render(ws, max_rows=last)
        wb.close()
    finally:
        out.unlink(missing_ok=True)
    return {"sheet": render, "header_row": hr, "start_row": sr, "last_row": last,
            "warnings": res.get("warnings", [])}


@router.post("/{tpl_id}/test-export")
def template_test_export(tpl_id: int, body: TrialIn, conn: sqlite3.Connection = Depends(get_db),
                         user=Depends(require_admin)):
    """Unduh .xlsx uji dari mapping yang sedang diedit (tanpa menyimpan)."""
    from fastapi.responses import FileResponse
    from starlette.background import BackgroundTask
    out, _res, _cfg = _trial_build(tpl_id, body, conn)
    return FileResponse(str(out), filename="Test_BAA.xlsx",
                        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        background=BackgroundTask(lambda: out.unlink(missing_ok=True)))
