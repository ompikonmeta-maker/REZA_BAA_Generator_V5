"""Import massal lokasi dari Excel (admin only).

Alur: unduh template (2 sheet) -> isi -> upload -> preview (dry-run, validasi)
-> commit (buat lokasi + inventory, kode urut otomatis).
"""
from __future__ import annotations

import datetime as _dt
import io
import json
import sqlite3

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from fastapi.responses import StreamingResponse

import re

from .. import db
from ..deps import audit, get_db, require_admin
from ..services import wilayah
from .locations import _format_code, _row_name, _wil_values, _WIL_SET, dup_key, norm_sid, sid_owner

router = APIRouter(prefix="/api/locations/import", tags=["locations-import"])

INV_COLS = ["nama_barang", "merk_type", "jumlah", "sn_tagging", "keterangan"]
# Kolom wilayah (bila fitur wilayah aktif di project): judul template + pengenal judul
WIL_COLS = [("kode", "Kode Wilayah", r"kode\s*wilayah|kode\s*desa"), ("desa", "Desa/Kelurahan", r"desa|kelurahan"), ("kec", "Kecamatan", r"kecamatan|\bkec\b"),
            ("kab", "Kab/Kota", r"kabupaten|\bkab\b|kota"), ("prov", "Provinsi", r"provinsi|propinsi|\bprov\b")]


# Kolom pemilik (opsional): username atau nama lengkap user project; kosong = admin pengimport
OWNER_RX = r"owner|pemilik|operator"
# Kolom SID (opsional): judul = label SID project, 'SID' atau 'Site ID'
SID_RX = r"^(sid|site\s*id)$"


def _sid_on(conn) -> bool:
    return bool(db.sid_config(conn).get("enabled", True))


def _sid_raw(v) -> str:
    """Sel Excel -> teks SID (angka 1.2E+11 / 123.0 dari Excel jadi digit utuh)."""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return _norm(v)


def _wil_on(conn) -> bool:
    return bool((getattr(conn, "project", None) or {}).get("wilayah_on", 1))
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _fields(conn):
    return db.get_setting(conn, "location_fields", []) or []


def _norm(v) -> str:
    return str(v).strip() if v is not None else ""


def _to_iso(v):
    """Kembalikan (iso, ok). Terima date/datetime atau string beberapa format."""
    if v is None or v == "":
        return "", True
    if isinstance(v, (_dt.datetime, _dt.date)):
        return v.strftime("%Y-%m-%d"), True
    s = str(v).strip()
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%Y/%m/%d"):
        try:
            return _dt.datetime.strptime(s, fmt).strftime("%Y-%m-%d"), True
        except ValueError:
            pass
    return s, False


@router.get("/template")
def template(conn: sqlite3.Connection = Depends(get_db), user=Depends(require_admin)):
    import openpyxl
    from openpyxl.utils import get_column_letter
    fields = _fields(conn)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Lokasi"
    wcols = [w[1] for w in WIL_COLS[1:]] + [WIL_COLS[0][1]] if _wil_on(conn) else []
    sid_lab = [db.sid_config(conn).get("label") or "SID"] if _sid_on(conn) else []
    ws.append(["ref"] + sid_lab + [(_norm(f.get("label")) or f.get("key")) for f in fields] + wcols + ["Owner"])
    ex = ["L1"] + [""] * len(sid_lab)      # SID: dari klien, kosong boleh
    for f in fields:
        if f.get("key") == "nama_lokasi":
            ex.append("Contoh: Desa Sukamaju")
        elif f.get("type") == "date":
            ex.append("2026-09-01")
        else:
            ex.append("")
    if wcols:
        ex += ["Cibodas", "Lembang", "Kab. Bandung Barat", "Jawa Barat", ""]
    ex.append("")                         # Owner: username / nama lengkap, kosong = admin pengimport
    ws.append(ex)
    wi = wb.create_sheet("Inventory")
    wi.append(["ref"] + INV_COLS)
    wi.append(["L1", "Access Point", "Starlink", "1", "SN-CONTOH", "Baik"])
    for w in (ws, wi):
        for c in range(1, w.max_column + 1):
            w.column_dimensions[get_column_letter(c)].width = 20
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf, media_type=XLSX_MIME,
        headers={"Content-Disposition": "attachment; filename=Template_Import_Lokasi.xlsx"},
    )


def _parse(conn: sqlite3.Connection, data: bytes) -> dict:
    import openpyxl
    fields = _fields(conn)
    label_to_field = {}
    for f in fields:
        label_to_field[_norm(f.get("label")).lower()] = f
        label_to_field[_norm(f.get("key")).lower()] = f
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    except Exception:
        raise HTTPException(400, "Can't read the file — make sure it's .xlsx")
    if "Lokasi" not in wb.sheetnames:
        raise HTTPException(400, "Sheet 'Lokasi' not found. Download the template first.")

    rows = list(wb["Lokasi"].iter_rows(values_only=True))
    if not rows:
        return {"locations": [], "summary": {"total": 0, "ready": 0, "items": 0, "errors": 0, "dup": 0}}
    header = [_norm(h) for h in rows[0]]
    col_field, ref_idx, wil_idx, own_idx, sid_idx = {}, None, {}, None, None
    sid_lab = (db.sid_config(conn).get("label") or "SID").strip().lower()
    for idx, h in enumerate(header):
        hl = h.lower()
        if hl == "ref":
            ref_idx = idx
        elif sid_idx is None and _sid_on(conn) and (hl == sid_lab or re.search(SID_RX, hl)):
            sid_idx = idx
        elif hl in label_to_field:
            col_field[idx] = label_to_field[hl]
        elif own_idx is None and re.search(OWNER_RX, hl):
            own_idx = idx
        elif _wil_on(conn):
            for key, _lab, rx in WIL_COLS:
                if key not in wil_idx.values() and re.search(rx, hl) and not (key == "kab" and "kode" in hl):
                    wil_idx[idx] = key
                    break

    # Inventory dikelompokkan per ref
    inv_by_ref: dict[str, list] = {}
    if "Inventory" in wb.sheetnames:
        irows = list(wb["Inventory"].iter_rows(values_only=True))
        if irows:
            ih = [_norm(h).lower() for h in irows[0]]
            cref = ih.index("ref") if "ref" in ih else None
            cmap = {c: (ih.index(c) if c in ih else None) for c in INV_COLS}
            for r in irows[1:]:
                if r is None:
                    continue
                ref = _norm(r[cref]) if (cref is not None and cref < len(r)) else ""
                item = {c: (_norm(r[cmap[c]]) if (cmap[c] is not None and cmap[c] < len(r)) else "") for c in INV_COLS}
                if not any(item.values()):
                    continue
                inv_by_ref.setdefault(ref, []).append(item)

    # Peta kanonik untuk auto-correct
    merks = db.get_setting(conn, "item_merks", {}) or {}
    kets = db.get_setting(conn, "inventory_keterangan", []) or []
    ket_map = {k.lower(): k for k in kets}

    def fix_merk(nama, val):
        for o in merks.get(_norm(nama), []):
            if o.lower() == val.lower():
                return o
        return val

    # Lokasi yang sudah ada: kunci kembar = nama + wilayah (nama sama di desa lain bukan duplikat)
    existing, by_name = {}, {}
    for r in conn.execute(
            """SELECT l.code, l.name, l.data_json, l.wil_desa, l.wil_kec, l.wil_kab,
                      COALESCE(NULLIF(TRIM(o.full_name),''), o.username, '—') AS owner_name
               FROM locations l LEFT JOIN users o ON o.id = l.owner_id WHERE l.deleted_at IS NULL ORDER BY l.id"""):
        nm = _row_name(r)
        info = {"code": r["code"], "desa": r["wil_desa"] or "", "owner": r["owner_name"]}
        k = dup_key(nm, r["wil_desa"], r["wil_kec"], r["wil_kab"])
        if k:
            existing.setdefault(k, info)
            by_name.setdefault(k.split("#")[0], info)
    # User project untuk kolom Owner (username / nama lengkap, tidak peka huruf besar/kecil)
    users = {}
    for u in conn.execute("SELECT id, username, full_name, role, active FROM pm_users"):
        for key in {(u["username"] or "").strip().lower(), (u["full_name"] or "").strip().lower()} - {""}:
            users.setdefault(key, u)

    def owner_of(raw: str) -> dict:
        if not raw:
            return {"status": "empty", "raw": ""}
        u = users.get(raw.lower())
        if not u:
            return {"status": "unknown", "raw": raw}
        name = (u["full_name"] or u["username"]).strip()
        if not u["active"]:
            return {"status": "inactive", "raw": raw, "name": name}
        if u["role"] == "viewer":
            return {"status": "viewer", "raw": raw, "name": name}
        return {"status": "ok", "raw": raw, "id": u["id"], "name": name}

    out, seen_ref, seen_key, seen_sid = [], {}, {}, {}
    for ri, r in enumerate(rows[1:], start=2):
        if r is None or all(_norm(c) == "" for c in r):
            continue
        data, errors = {}, []
        for idx, f in col_field.items():
            val = r[idx] if idx < len(r) else None
            if f.get("type") == "date":
                iso, ok = _to_iso(val)
                if not ok:
                    errors.append(f"Tanggal '{f.get('label')}' tidak valid")
                data[f.get("key")] = iso
            else:
                data[f.get("key")] = _norm(val)
        nama = _norm(data.get("nama_lokasi"))
        if not nama:
            errors.append("Nama lokasi wajib")
        ref = _norm(r[ref_idx]) if (ref_idx is not None and ref_idx < len(r)) else ""
        if ref:
            if ref in seen_ref:
                errors.append(f"ref '{ref}' dipakai lebih dari sekali")
            seen_ref[ref] = ri
        inv = []
        for it in (inv_by_ref.get(ref, []) if ref else []):
            it = dict(it)
            it["merk_type"] = fix_merk(it.get("nama_barang", ""), it.get("merk_type", ""))
            it["keterangan"] = ket_map.get(it.get("keterangan", "").lower(), it.get("keterangan", ""))
            inv.append(it)
        wil = None
        if wil_idx:
            wv = {k: _norm(r[i]) if i < len(r) else "" for i, k in wil_idx.items()}
            wil = wilayah.match(wv.get("desa", ""), wv.get("kec", ""), wv.get("kab", ""), wv.get("prov", ""),
                                wv.get("kode", ""))
            wil["raw"] = {k: wv.get(k, "") for k in ("desa", "kec", "kab", "prov")}
        # kembar: ke lokasi project & ke baris sebelumnya di file yang sama (nama + wilayah)
        ww = (wil.get("wil") if wil and wil.get("status") in ("ok", "desa_manual") else None) or (wil or {}).get("raw") or {}
        k = dup_key(nama, ww.get("desa", ""), ww.get("kec", ""), ww.get("kab", ""))
        dup_of = existing.get(k) if k else None
        dup_file = seen_key.get(k) if (k and not dup_of) else None
        if k and not errors:
            seen_key.setdefault(k, ri)
        same_name = None
        if k and not dup_of and not dup_file:
            same_name = by_name.get(k.split("#")[0])
        own = owner_of(_norm(r[own_idx]) if (own_idx is not None and own_idx < len(r)) else "")
        # SID: dirapikan (kapital bila aktif). Sudah ada di project -> baris dilewati (data lama tidak ditimpa);
        # dobel di file -> baris berikutnya error (SID harus unik)
        sid, sid_of, sid_in = "", None, ""
        if sid_idx is not None:
            sid_in = _sid_raw(r[sid_idx]) if sid_idx < len(r) else ""
            try:
                sid = norm_sid(conn, sid_in)
            except HTTPException:
                lab = db.sid_config(conn).get("label") or "SID"
                errors.append(f"{lab} '{sid_in}' tidak valid (huruf & angka, maks {db.sid_config(conn).get('max_len') or 20})")
            if sid:
                o = sid_owner(conn, sid)
                if o:
                    sid_of = o["code"]
                elif sid.upper() in seen_sid:
                    errors.append(f"{db.sid_config(conn).get('label') or 'SID'} {sid} dobel di file (baris {seen_sid[sid.upper()]})")
                else:
                    seen_sid[sid.upper()] = ri
        out.append({"row": ri, "nama": nama, "data": data, "inventory": inv,
                    "errors": errors, "dup": bool(dup_of or dup_file), "dup_of": dup_of, "dup_file": dup_file,
                    "same_name": same_name, "wil": wil, "owner": own, "sid": sid, "sid_in": sid_in, "sid_of": sid_of})

    ok_rows = [x for x in out if not x["errors"]]
    owners: dict[str, int] = {}
    for x in ok_rows:
        if x["owner"]["status"] == "ok":
            owners[x["owner"]["name"]] = owners.get(x["owner"]["name"], 0) + 1
    return {"locations": out, "summary": {
        "total": len(out),
        "ready": sum(1 for x in ok_rows if not x["sid_of"]),
        "items": sum(len(x["inventory"]) for x in ok_rows),
        "errors": sum(1 for x in out if x["errors"]),
        "dup": sum(1 for x in ok_rows if x["dup"] and not x["sid_of"]),
        "sid_col": sid_idx is not None,
        "sid_n": sum(1 for x in ok_rows if x["sid"] and not x["sid_of"]),
        "sid_skip": sum(1 for x in ok_rows if x["sid_of"]),
        "wil_cols": bool(wil_idx),
        "wil_ok": sum(1 for x in ok_rows if x["wil"] and x["wil"]["status"] == "ok"),
        "wil_manual": sum(1 for x in ok_rows if x["wil"] and x["wil"]["status"] == "desa_manual"),
        "wil_check": sum(1 for x in ok_rows if x["wil"] and x["wil"]["status"] in ("ambiguous", "notfound")),
        "own_col": own_idx is not None, "owners": owners,
        "own_me": sum(1 for x in ok_rows if x["owner"]["status"] == "empty"),
        "own_check": sum(1 for x in ok_rows if x["owner"]["status"] in ("unknown", "viewer", "inactive")),
    }}


@router.post("/preview")
async def preview(file: UploadFile = File(...), conn: sqlite3.Connection = Depends(get_db),
                  user=Depends(require_admin)):
    return _parse(conn, await file.read())


@router.post("/commit")
async def commit(file: UploadFile = File(...), skip_dup: bool = Form(True), wil_choices: str = Form("{}"),
                 owner_choices: str = Form("{}"),
                 conn: sqlite3.Connection = Depends(get_db), user=Depends(require_admin)):
    """wil_choices: {baris: {"kode": desa resmi} | {"manual": 1} | {}}. Baris wilayah yang perlu dicek
    tanpa pilihan disimpan TANPA wilayah (tidak ada penyimpanan diam-diam)."""
    res = _parse(conn, await file.read())
    try:
        choices = {str(k): v for k, v in (json.loads(wil_choices or "{}") or {}).items()}
    except Exception:
        choices = {}
    try:
        own_ch = {str(k): v for k, v in (json.loads(owner_choices or "{}") or {}).items()}
    except Exception:
        own_ch = {}
    # pemilik yang boleh dipilih: user project aktif, bukan viewer
    can_own = {r["id"] for r in conn.execute("SELECT id FROM pm_users WHERE active=1 AND role<>'viewer'")}
    created = skipped = 0
    now = db.now_iso()
    for loc in res["locations"]:
        if loc["errors"] or loc.get("sid_of") or (loc["dup"] and skip_dup):
            skipped += 1
            continue
        # pemilik: pilihan admin di preview > kolom Owner yang cocok > admin pengimport; created_by tetap admin
        ch = own_ch.get(str(loc["row"]))
        owner = int(ch) if str(ch or "").isdigit() and int(ch) in can_own else (
            loc["owner"]["id"] if loc["owner"]["status"] == "ok" and ch is None else user["id"])
        cur = conn.execute(
            "INSERT INTO locations(code,name,data_json,status,created_by,owner_id,created_at,updated_at,sid) "
            "VALUES(?,?,?,?,?,?,?,?,?)",
            ("", loc["nama"], json.dumps(loc["data"], ensure_ascii=False), "draft",
             user["id"], owner, now, now, loc.get("sid") or None),
        )
        lid = cur.lastrowid
        conn.execute("UPDATE locations SET code=?, modified_at=?, modified_by=? WHERE id=?",
                     (_format_code(conn, lid), now, user["id"], lid))
        wv = _import_wil(loc.get("wil"), choices.get(str(loc["row"])) or {})
        if wv:
            conn.execute(f"UPDATE locations SET {_WIL_SET} WHERE id=?", (*wv, lid))
        db.log_activity(conn, lid, user["id"], "create", "import", now)
        for i, it in enumerate(loc["inventory"]):
            conn.execute(
                "INSERT INTO inventory_items(location_id,nama_barang,merk_type,jumlah,"
                "sn_tagging,keterangan,sort_order) VALUES(?,?,?,?,?,?,?)",
                (lid, it.get("nama_barang", ""), it.get("merk_type", ""), it.get("jumlah", ""),
                 it.get("sn_tagging", ""), it.get("keterangan", ""), i),
            )
        created += 1
    conn.commit()
    audit(conn, user, "import", "location", "", f"{created} lokasi, {skipped} dilewati")
    return {"ok": True, "created": created, "skipped": skipped}


def _import_wil(w: dict | None, ch: dict):
    """Kolom wilayah untuk satu baris import (atau None = tanpa wilayah)."""
    if not w or w.get("status") == "empty":
        return None
    raw = w.get("raw") or {}
    try:
        if ch.get("kode"):
            return _wil_values({"mode": "official", "kode": ch["kode"], **raw})
        if ch.get("manual") and all(raw.get(k) for k in ("desa", "kec", "kab", "prov")):
            return _wil_values({"mode": "manual", **raw})
        if ch.get("empty"):
            return None
        if w.get("status") in ("ok", "desa_manual"):
            return _wil_values(w["wil"])
    except HTTPException:
        return None
    return None


# ================= Update SID massal (lokasi yang sudah ada) =================
# Excel: Kode lokasi -> SID. Hanya MENGISI SID yang masih kosong; SID yang sudah terisi tidak ditimpa
# (perubahan tetap lewat admin per lokasi / permintaan revisi). Aturan sama dengan isian manual.
CODE_RX = r"^(kode|code|kode lokasi|location code|kode baa)$"


@router.get("/sid-template")
def sid_template(conn: sqlite3.Connection = Depends(get_db), user=Depends(require_admin)):
    """Daftar semua lokasi aktif (kode, nama, desa, SID saat ini) — admin tinggal mengisi kolom SID."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    lab = db.sid_config(conn).get("label") or "SID"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "SID"
    ws.append(["Kode", "Nama Lokasi", "Desa/Kelurahan", lab])
    for c in ws[1]:
        c.font = Font(bold=True)
    grey = PatternFill("solid", fgColor="EEEEEE")
    for r in conn.execute("SELECT code, name, data_json, wil_desa, sid FROM locations WHERE deleted_at IS NULL ORDER BY id"):
        ws.append([r["code"], _row_name(r), r["wil_desa"] or "", r["sid"] or ""])
        if r["sid"]:
            ws.cell(ws.max_row, 4).fill = grey            # sudah terisi: tidak akan ditimpa
    for col, w in zip("ABCD", (14, 44, 22, 22)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    pre = (getattr(conn, "project", None) or {}).get("prefix") or "Lokasi"
    return StreamingResponse(buf, media_type=XLSX_MIME,
                             headers={"Content-Disposition": f"attachment; filename=Update_{lab}_{pre}.xlsx"})


def _sid_parse(conn, data: bytes) -> dict:
    import openpyxl
    lab = db.sid_config(conn).get("label") or "SID"
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    except Exception:
        raise HTTPException(400, "Can't read the file — make sure it's .xlsx")
    ws = wb["SID"] if "SID" in wb.sheetnames else wb.active
    rows = list(ws.iter_rows(values_only=True))
    header = [_norm(h).lower() for h in (rows[0] if rows else [])]
    ci = next((i for i, h in enumerate(header) if re.search(CODE_RX, h)), None)
    si = next((i for i, h in enumerate(header) if h == lab.lower() or re.search(SID_RX, h)), None)
    if ci is None or si is None:
        raise HTTPException(400, f"The file needs a 'Kode' column and a '{lab}' column — download the list first")
    locs, by_old = {}, {}
    for r in conn.execute("SELECT id, code, old_codes, name, data_json, sid FROM locations WHERE deleted_at IS NULL"):
        locs[r["code"].upper()] = r
        for o in (r["old_codes"] or "").split():
            by_old.setdefault(o.upper(), r)
    out, seen_code, seen_sid = [], {}, {}
    for ri, r in enumerate(rows[1:], start=2):
        if r is None or all(_norm(c) == "" for c in r):
            continue
        code = _norm(r[ci]) if ci < len(r) else ""
        sid_in = _sid_raw(r[si]) if si < len(r) else ""
        loc = locs.get(code.upper()) or by_old.get(code.upper())
        row = {"row": ri, "code": code, "loc_code": loc["code"] if loc else "", "nama": _row_name(loc) if loc else "",
               "sid_in": sid_in, "sid": "", "cur": (loc["sid"] or "") if loc else "", "status": "ok", "note": ""}
        out.append(row)
        if not sid_in:
            row["status"] = "empty"
            continue
        try:
            row["sid"] = norm_sid(conn, sid_in)
        except HTTPException as e:
            row["status"], row["note"] = "invalid", str(e.detail)
            continue
        if not loc:
            row["status"] = "notfound"
        elif loc["id"] in seen_code:
            row["status"], row["note"] = "dup_file", f"same location as row {seen_code[loc['id']]}"
        elif row["cur"] and row["cur"].upper() == row["sid"].upper():
            row["status"] = "same"
        elif row["cur"]:
            row["status"] = "filled"
        elif row["sid"].upper() in seen_sid:
            row["status"], row["note"] = "dup_file", f"same {lab} as row {seen_sid[row['sid'].upper()]}"
        elif (o := sid_owner(conn, row["sid"], loc["id"])):
            row["status"], row["note"] = "used", o["code"]
        if loc:
            seen_code.setdefault(loc["id"], ri)
        if row["status"] == "ok":
            seen_sid[row["sid"].upper()] = ri
            row["id"] = loc["id"]
    cnt = {}
    for x in out:
        cnt[x["status"]] = cnt.get(x["status"], 0) + 1
    return {"rows": out, "summary": {"total": len(out), **cnt, "ok": cnt.get("ok", 0)}}


@router.post("/sid-preview")
async def sid_preview(file: UploadFile = File(...), conn: sqlite3.Connection = Depends(get_db),
                      user=Depends(require_admin)):
    res = _sid_parse(conn, await file.read())
    for x in res["rows"]:
        x.pop("id", None)
    return res


@router.post("/sid-commit")
async def sid_commit(file: UploadFile = File(...), conn: sqlite3.Connection = Depends(get_db),
                     user=Depends(require_admin)):
    res = _sid_parse(conn, await file.read())
    now, n = db.now_iso(), 0
    for x in res["rows"]:
        if x["status"] != "ok":
            continue
        # tetap hanya mengisi yang kosong (jaga-jaga diisi orang lain sejak preview)
        cur = conn.execute("UPDATE locations SET sid=?, updated_at=? WHERE id=? AND (sid IS NULL OR sid='')",
                           (x["sid"], now, x["id"]))
        if cur.rowcount:
            db.log_activity(conn, x["id"], user["id"], "sid", f"— -> {x['sid']} (import)", now)
            n += 1
    conn.commit()
    audit(conn, user, "import_sid", "location", "", f"{n} SID diisi")
    return {"ok": True, "updated": n, "skipped": res["summary"]["total"] - n}
